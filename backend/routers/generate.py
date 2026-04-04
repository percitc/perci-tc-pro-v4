"""Router /generate — con exportacion PDF y DOCX"""
from __future__ import annotations
import io, json
from fastapi import APIRouter, HTTPException
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel
from services.rag_service import rag_service
from services.generators import generate, podcast_audio, tts

router = APIRouter()


class GenerateRequest(BaseModel):
    session_id: str
    output_type: str
    query: str | None = None
    title: str | None = None
    num_items: int | None = None


class TTSRequest(BaseModel):
    text: str
    voice: str = "onyx"


class ExportRequest(BaseModel):
    content: dict
    output_type: str
    format: str  # "pdf" o "docx"


# ── Generar contenido ──────────────────────────────────────────────────────────

@router.post("/content")
async def gen_content(req: GenerateRequest):
    query  = req.query or "resume y explica el contenido completo"
    chunks = await rag_service.retrieve(req.session_id, query)
    if not chunks:
        raise HTTPException(404, "Sin documentos en esta sesion")
    context = "\n\n---\n\n".join(c["text"] for c in chunks)
    kwargs: dict = {}
    if req.title: kwargs["title"] = req.title
    if req.num_items:
        if req.output_type in ("flashcards", "quiz"): kwargs["num"] = req.num_items
        elif req.output_type == "slides":             kwargs["num_slides"] = req.num_items
    try:
        return JSONResponse(await generate(req.output_type, context, **kwargs))
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))


# ── Exportar PDF ───────────────────────────────────────────────────────────────

def _build_pdf(content: dict, output_type: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()

    # Estilos personalizados
    title_style = ParagraphStyle('Title2', parent=styles['Title'],
                                 fontSize=18, textColor=colors.HexColor('#1e40af'), spaceAfter=12)
    h2_style    = ParagraphStyle('H2', parent=styles['Heading2'],
                                 fontSize=13, textColor=colors.HexColor('#1e40af'), spaceAfter=8)
    body_style  = ParagraphStyle('Body2', parent=styles['Normal'],
                                 fontSize=10, leading=14, spaceAfter=6)
    label_style = ParagraphStyle('Label', parent=styles['Normal'],
                                 fontSize=9, textColor=colors.HexColor('#6b7280'))

    story = []

    # ── Resumen ──────────────────────────────────────────────────────────────
    if output_type == "summary":
        story.append(Paragraph("Resumen Ejecutivo", title_style))
        story.append(Spacer(1, 0.3*cm))
        text = content if isinstance(content, str) else str(content)
        for line in text.split("\n"):
            if line.strip():
                story.append(Paragraph(line.strip(), body_style))

    # ── Flashcards ───────────────────────────────────────────────────────────
    elif output_type == "flashcards":
        story.append(Paragraph("Tarjetas Didácticas", title_style))
        story.append(Spacer(1, 0.3*cm))
        tarjetas = content.get("tarjetas", content) if isinstance(content, dict) else []
        for t in tarjetas:
            data = [
                [Paragraph(f"<b>#{t.get('id','')}</b> {t.get('categoria','')}", label_style),
                 Paragraph(t.get('dificultad',''), label_style)],
                [Paragraph(f"<b>P:</b> {t.get('pregunta','')}", body_style), ""],
                [Paragraph(f"<b>R:</b> {t.get('respuesta','')}", body_style), ""],
            ]
            tbl = Table(data, colWidths=[13*cm, 3*cm])
            tbl.setStyle(TableStyle([
                ('BOX',        (0,0), (-1,-1), 0.5, colors.HexColor('#e5e7eb')),
                ('BACKGROUND', (0,0), (-1,0),  colors.HexColor('#eff6ff')),
                ('SPAN',       (0,1), (1,1)),
                ('SPAN',       (0,2), (1,2)),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f9fafb')]),
                ('TOPPADDING',  (0,0), (-1,-1), 6),
                ('BOTTOMPADDING',(0,0),(-1,-1), 6),
                ('LEFTPADDING', (0,0), (-1,-1), 8),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 0.2*cm))

    # ── Cuestionario ─────────────────────────────────────────────────────────
    elif output_type == "quiz":
        story.append(Paragraph("Cuestionario", title_style))
        story.append(Spacer(1, 0.3*cm))
        preguntas = content.get("cuestionario", content) if isinstance(content, dict) else []
        for q in preguntas:
            story.append(Paragraph(f"<b>{q.get('id','')}. {q.get('pregunta','')}</b>", body_style))
            for op in q.get("opciones", []):
                story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;{op}", body_style))
            resp = q.get('respuesta_correcta','')
            story.append(Paragraph(f"<font color='#16a34a'><b>✓ Respuesta: {resp}</b></font>", label_style))
            story.append(Paragraph(f"<i>{q.get('explicacion','')}</i>", label_style))
            story.append(Spacer(1, 0.3*cm))

    # ── Slides ────────────────────────────────────────────────────────────────
    elif output_type == "slides":
        titulo = content.get("titulo_presentacion", "Presentación") if isinstance(content, dict) else "Presentación"
        story.append(Paragraph(titulo, title_style))
        story.append(Spacer(1, 0.3*cm))
        for slide in content.get("slides", []):
            story.append(Paragraph(f"{slide.get('emoji','')} {slide.get('titulo','')}", h2_style))
            if slide.get("subtitulo"):
                story.append(Paragraph(slide["subtitulo"], label_style))
            for punto in slide.get("puntos", []):
                story.append(Paragraph(f"• {punto}", body_style))
            if slide.get("nota_orador"):
                story.append(Paragraph(f"<i>Nota: {slide['nota_orador']}</i>", label_style))
            story.append(Spacer(1, 0.4*cm))

    # ── Curso ─────────────────────────────────────────────────────────────────
    elif output_type == "course":
        titulo = content.get("titulo_curso", "Curso") if isinstance(content, dict) else "Curso"
        story.append(Paragraph(titulo, title_style))
        story.append(Paragraph(content.get("descripcion",""), body_style))
        story.append(Paragraph(f"Duración: {content.get('duracion_estimada','')} | Nivel: {content.get('nivel','')}", label_style))
        story.append(Spacer(1, 0.3*cm))
        story.append(Paragraph("Objetivos", h2_style))
        for obj in content.get("objetivos", []):
            story.append(Paragraph(f"• {obj}", body_style))
        story.append(Spacer(1, 0.2*cm))
        for mod in content.get("modulos", []):
            story.append(Paragraph(f"Módulo {mod.get('numero','')} — {mod.get('titulo','')}", h2_style))
            story.append(Paragraph(mod.get("descripcion",""), body_style))
            for tema in mod.get("temas", []):
                story.append(Paragraph(f"<b>{tema.get('titulo','')}</b>", body_style))
                story.append(Paragraph(tema.get("contenido",""), body_style))
            story.append(Spacer(1, 0.2*cm))

    # ── Infografía ────────────────────────────────────────────────────────────
    elif output_type == "infographic":
        story.append(Paragraph(content.get("titulo","Infografía"), title_style))
        story.append(Paragraph(content.get("subtitulo",""), body_style))
        story.append(Spacer(1, 0.3*cm))
        for sec in content.get("secciones", []):
            story.append(Paragraph(f"{sec.get('icono','')} {sec.get('titulo','')}", h2_style))
            story.append(Paragraph(sec.get("descripcion",""), body_style))
            for dato in sec.get("datos", []):
                story.append(Paragraph(f"• {dato}", body_style))
        if content.get("conclusion"):
            story.append(Spacer(1, 0.3*cm))
            story.append(Paragraph(f"💡 {content['conclusion']}", body_style))

    # ── Guión ─────────────────────────────────────────────────────────────────
    elif output_type == "audio_script":
        story.append(Paragraph("Guión de Podcast Educativo", title_style))
        story.append(Spacer(1, 0.3*cm))
        text = content if isinstance(content, str) else str(content)
        for line in text.split("\n"):
            if line.strip():
                if line.startswith("[PROFESOR]"):
                    story.append(Paragraph(f"<b><font color='#1e40af'>{line}</font></b>", body_style))
                elif line.startswith("[ASISTENTE]"):
                    story.append(Paragraph(f"<b><font color='#15803d'>{line}</font></b>", body_style))
                else:
                    story.append(Paragraph(line, body_style))

    else:
        story.append(Paragraph("Contenido", title_style))
        story.append(Paragraph(str(content), body_style))

    doc.build(story)
    return buf.getvalue()


# ── Exportar DOCX ──────────────────────────────────────────────────────────────

def _build_docx(content: dict, output_type: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = Document()
    BLUE = RGBColor(0x1e, 0x40, 0xaf)
    GREEN = RGBColor(0x15, 0x80, 0x3d)

    def add_title(text):
        p = document.add_heading(text, level=1)
        p.runs[0].font.color.rgb = BLUE

    def add_h2(text):
        p = document.add_heading(text, level=2)
        if p.runs: p.runs[0].font.color.rgb = BLUE

    def add_body(text):
        p = document.add_paragraph(text)
        p.runs[0].font.size = Pt(11) if p.runs else None

    # ── Por tipo ──────────────────────────────────────────────────────────────
    if output_type == "summary":
        add_title("Resumen Ejecutivo")
        text = content if isinstance(content, str) else str(content)
        for line in text.split("\n"):
            if line.strip():
                document.add_paragraph(line.strip())

    elif output_type == "flashcards":
        add_title("Tarjetas Didácticas")
        tarjetas = content.get("tarjetas", []) if isinstance(content, dict) else []
        for t in tarjetas:
            p = document.add_paragraph()
            p.add_run(f"#{t.get('id','')} — {t.get('categoria','')} ({t.get('dificultad','')})").bold = True
            document.add_paragraph(f"Pregunta: {t.get('pregunta','')}")
            p2 = document.add_paragraph(f"Respuesta: {t.get('respuesta','')}")
            p2.runs[0].font.color.rgb = GREEN
            document.add_paragraph("─" * 60)

    elif output_type == "quiz":
        add_title("Cuestionario")
        preguntas = content.get("cuestionario", []) if isinstance(content, dict) else []
        for q in preguntas:
            p = document.add_paragraph()
            p.add_run(f"{q.get('id','')}. {q.get('pregunta','')}").bold = True
            for op in q.get("opciones", []):
                document.add_paragraph(f"   {op}", style="List Bullet")
            p2 = document.add_paragraph(f"✓ Respuesta correcta: {q.get('respuesta_correcta','')}")
            p2.runs[0].font.color.rgb = GREEN
            document.add_paragraph(f"Explicación: {q.get('explicacion','')}")
            document.add_paragraph("")

    elif output_type == "slides":
        titulo = content.get("titulo_presentacion","Presentación") if isinstance(content, dict) else "Presentación"
        add_title(titulo)
        for slide in content.get("slides", []):
            add_h2(f"{slide.get('emoji','')} {slide.get('titulo','')}")
            if slide.get("subtitulo"):
                document.add_paragraph(slide["subtitulo"])
            for punto in slide.get("puntos", []):
                document.add_paragraph(punto, style="List Bullet")
            if slide.get("nota_orador"):
                p = document.add_paragraph(f"Nota del orador: {slide['nota_orador']}")
                p.runs[0].italic = True

    elif output_type == "course":
        add_title(content.get("titulo_curso","Curso"))
        document.add_paragraph(content.get("descripcion",""))
        document.add_paragraph(f"Duración: {content.get('duracion_estimada','')} | Nivel: {content.get('nivel','')}")
        add_h2("Objetivos")
        for obj in content.get("objetivos", []):
            document.add_paragraph(obj, style="List Bullet")
        for mod in content.get("modulos", []):
            add_h2(f"Módulo {mod.get('numero','')} — {mod.get('titulo','')}")
            document.add_paragraph(mod.get("descripcion",""))
            for tema in mod.get("temas", []):
                p = document.add_paragraph()
                p.add_run(tema.get("titulo","")).bold = True
                document.add_paragraph(tema.get("contenido",""))

    elif output_type == "infographic":
        add_title(content.get("titulo","Infografía"))
        document.add_paragraph(content.get("subtitulo",""))
        for sec in content.get("secciones", []):
            add_h2(f"{sec.get('icono','')} {sec.get('titulo','')}")
            document.add_paragraph(sec.get("descripcion",""))
            for dato in sec.get("datos", []):
                document.add_paragraph(dato, style="List Bullet")
        if content.get("conclusion"):
            p = document.add_paragraph(f"Conclusión: {content['conclusion']}")
            p.runs[0].bold = True

    elif output_type == "audio_script":
        add_title("Guión de Podcast Educativo")
        text = content if isinstance(content, str) else str(content)
        for line in text.split("\n"):
            if line.strip():
                p = document.add_paragraph()
                run = p.add_run(line)
                if line.startswith("[PROFESOR]"):
                    run.font.color.rgb = BLUE
                    run.bold = True
                elif line.startswith("[ASISTENTE]"):
                    run.font.color.rgb = GREEN
                    run.bold = True

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ── Endpoint exportar ──────────────────────────────────────────────────────────

@router.post("/export")
async def export_content(req: ExportRequest):
    """Exporta contenido generado a PDF o DOCX."""
    content = req.content.get("content", req.content)
    fmt     = req.format.lower()

    try:
        if fmt == "pdf":
            data     = _build_pdf(content, req.output_type)
            filename = f"{req.output_type}.pdf"
            media    = "application/pdf"
        elif fmt == "docx":
            data     = _build_docx(content, req.output_type)
            filename = f"{req.output_type}.docx"
            media    = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            raise HTTPException(400, "Formato no soportado. Usa 'pdf' o 'docx'")

        return StreamingResponse(
            iter([data]),
            media_type=media,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Error al exportar: {e}")


# ── Audio ──────────────────────────────────────────────────────────────────────

@router.post("/audio/tts")
async def gen_tts(req: TTSRequest):
    try:
        audio = await tts(req.text, req.voice)
        return StreamingResponse(iter([audio]), media_type="audio/mpeg",
                                 headers={"Content-Disposition": 'attachment; filename="audio.mp3"'})
    except Exception as e:
        raise HTTPException(500, str(e))


@router.post("/audio/podcast")
async def gen_podcast(req: GenerateRequest):
    from services.generators import generate as gen
    chunks  = await rag_service.retrieve(req.session_id, req.query or "explica el contenido")
    if not chunks: raise HTTPException(404, "Sin documentos")
    context = "\n\n".join(c["text"] for c in chunks)
    script  = (await gen("audio_script", context))["content"]
    audio   = await podcast_audio(script)
    return StreamingResponse(iter([audio]), media_type="audio/mpeg",
                             headers={"Content-Disposition": 'attachment; filename="podcast.mp3"'})
