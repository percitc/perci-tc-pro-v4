"""PERCI TC PRO AI — Procesador de Archivos Multiformat"""
from __future__ import annotations
import os, io
from pathlib import Path
from dotenv import load_dotenv

load_dotenv()
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", "./data/uploads"))
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


async def _ocr_bytes(img_bytes: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(io.BytesIO(img_bytes)), lang="spa+eng")
    except Exception:
        from services.ai import ai_router
        import base64
        b64 = base64.b64encode(img_bytes).decode()
        r = await ai_router.generate_response(
            prompt="Extrae TODO el texto de esta imagen. Responde SOLO con el texto.",
            options={"max_tokens": 4000,
                     "messages": [{"role": "user", "content": [
                         {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}},
                         {"type": "text", "text": "Extrae todo el texto."}
                     ]}]},
        )
        return r.content


async def extract_pdf(p: Path) -> str:
    import fitz
    doc   = fitz.open(str(p))
    pages = []
    for page in doc:
        t = page.get_text("text")
        if t.strip():
            pages.append(t)
        else:
            pix = page.get_pixmap(dpi=200)
            pages.append(await _ocr_bytes(pix.tobytes("png")))
    doc.close()
    return "\n\n".join(pages)


async def extract_docx(p: Path) -> str:
    from docx import Document
    doc   = Document(str(p))
    parts = [pa.text for pa in doc.paragraphs if pa.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


async def extract_pptx(p: Path) -> str:
    from pptx import Presentation
    prs   = Presentation(str(p))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if texts:
            slides.append(f"--- Diapositiva {i} ---\n" + "\n".join(texts))
    return "\n\n".join(slides)


async def extract_txt(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="ignore")


async def extract_csv(p: Path) -> str:
    import pandas as pd
    try:
        df = pd.read_excel(str(p)) if p.suffix.lower() in [".xlsx",".xls"] else pd.read_csv(str(p), errors="ignore")
        return df.to_string(index=False, max_rows=500)
    except Exception as e:
        return f"Error leyendo {p.name}: {e}"


async def extract_image(p: Path) -> str:
    return await _ocr_bytes(p.read_bytes())


async def extract_audio(p: Path) -> str:
    from services.ai import ai_router
    r = await ai_router.transcribe_audio(str(p))
    return r.text


async def extract_video(p: Path) -> str:
    import ffmpeg
    audio_path = p.with_suffix(".wav")
    try:
        ffmpeg.input(str(p)).output(str(audio_path), acodec="pcm_s16le", ac=1, ar="16000").overwrite_output().run(quiet=True)
        return await extract_audio(audio_path)
    finally:
        if audio_path.exists(): audio_path.unlink()


async def extract_url(url: str) -> str:
    if "youtube.com" in url or "youtu.be" in url:
        return await _youtube(url)
    import httpx
    from html.parser import HTMLParser

    class _P(HTMLParser):
        def __init__(self): super().__init__(); self.parts=[]; self._skip=False
        def handle_starttag(self, t, a):
            if t in ("script","style","nav","footer","header"): self._skip=True
        def handle_endtag(self, t):
            if t in ("script","style","nav","footer","header"): self._skip=False
        def handle_data(self, d):
            if not self._skip and d.strip(): self.parts.append(d.strip())

    async with httpx.AsyncClient(follow_redirects=True, timeout=30) as h:
        resp = await h.get(url)
    p = _P(); p.feed(resp.text)
    return "\n".join(p.parts)


async def _youtube(url: str) -> str:
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
        import re
        vid = re.search(r"(?:v=|youtu\.be/)([^&?/]+)", url)
        if not vid: return "URL de YouTube invalida"
        t = YouTubeTranscriptApi.get_transcript(vid.group(1), languages=["es","en"])
        return " ".join(x["text"] for x in t)
    except Exception as e:
        return f"No se pudo obtener transcripcion de YouTube: {e}"


EXTRACTORS = {
    ".pdf": extract_pdf, ".docx": extract_docx, ".doc": extract_docx,
    ".pptx": extract_pptx, ".ppt": extract_pptx,
    ".txt": extract_txt, ".md": extract_txt,
    ".csv": extract_csv, ".xlsx": extract_csv, ".xls": extract_csv,
    ".png": extract_image, ".jpg": extract_image, ".jpeg": extract_image, ".webp": extract_image,
    ".mp3": extract_audio, ".wav": extract_audio, ".m4a": extract_audio,
    ".mp4": extract_video, ".mov": extract_video, ".avi": extract_video,
}


async def process_file(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    fn  = EXTRACTORS.get(ext)
    if not fn: raise ValueError(f"Extension no soportada: {ext}")
    return await fn(file_path)


async def process_url(url: str) -> str:
    return await extract_url(url)
